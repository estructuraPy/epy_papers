"""Tests for ``epy_papers._core._packaging.make_reference_docx``.

Dev-only release tooling; nothing exercised it before. ``_REF_DIR`` is
always monkeypatched to a ``tmp_path``: the real one resolves to
``src/epy_papers/_config/_assets/reference_docx`` -- INSIDE ``src/`` --
where the real bundled ``submission.docx`` / ``submission_lineno.docx``
live. A real ``build()`` call would overwrite them, which the task's
"tests only, never touch src/" rule forbids.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING

from epy_papers._core._packaging import make_reference_docx as mrd


def test_build_writes_us_letter_geometry_and_body_font(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setattr(mrd, "_REF_DIR", tmp_path)
    out = mrd.build(line_numbered=False)

    assert out == tmp_path / "submission.docx"
    assert out.exists()
    doc = Document(str(out))
    section = doc.sections[0]
    assert section.page_width == mrd.Inches(8.5)
    assert section.page_height == mrd.Inches(11)
    assert section.top_margin == mrd.Inches(1)
    normal = doc.styles["Normal"]
    assert normal.font.name == mrd.BODY_FONT
    assert normal.font.size == mrd.BODY_SIZE
    assert (
        normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE
    )


def test_build_seeds_every_heading_style_and_a_body_paragraph(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setattr(mrd, "_REF_DIR", tmp_path)
    out = mrd.build(line_numbered=False)
    doc = Document(str(out))
    styles_used = {p.style.name for p in doc.paragraphs}
    assert {"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"} <= (
        styles_used
    )
    assert any(
        "Times New Roman 12 pt" in p.text for p in doc.paragraphs
    )


def test_build_line_numbered_writes_a_different_file_with_lnnumtype(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    from docx.oxml.ns import qn

    monkeypatch.setattr(mrd, "_REF_DIR", tmp_path)
    out = mrd.build(line_numbered=True)

    assert out == tmp_path / "submission_lineno.docx"
    doc = Document(str(out))
    sect_pr = doc.sections[0]._sectPr
    ln = sect_pr.find(qn("w:lnNumType"))
    assert ln is not None
    assert ln.get(qn("w:countBy")) == "1"
    assert ln.get(qn("w:restart")) == "continuous"


def test_build_creates_the_output_directory_if_missing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    target_dir = tmp_path / "not_yet_created"
    monkeypatch.setattr(mrd, "_REF_DIR", target_dir)
    out = mrd.build(line_numbered=False)
    assert out.exists()


def test_add_line_numbering_inserts_before_existing_cols_element(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """When a ``<w:cols>`` element is already present, the line-number
    element must be inserted before it (required element ordering in the
    OOXML schema), not just appended at the end.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    monkeypatch.setattr(mrd, "_REF_DIR", tmp_path)
    doc = Document()
    section = doc.sections[0]
    cols = OxmlElement("w:cols")
    section._sectPr.append(cols)

    mrd._add_line_numbering(section)

    sect_pr = section._sectPr
    children = list(sect_pr)
    ln_index = children.index(sect_pr.find(qn("w:lnNumType")))
    cols_index = children.index(cols)
    assert ln_index < cols_index
